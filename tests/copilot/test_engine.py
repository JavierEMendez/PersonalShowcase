"""The multifamily engine on Sawyer Bend: headline ranges, footing, cases, stresses, fixture."""

import io
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from core.copilot.engine import run
from core.copilot.excel import export_workbook
from core.copilot.inputs import CopilotInputs
from core.copilot.sensitivity import at_price, stress_table
from core.copilot.summary import CopilotOutputs
from tests.copilot.conftest import load_case

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "sawyer_bend.json"


@pytest.fixture(scope="module")
def base(sawyer_base: CopilotInputs) -> CopilotOutputs:
    return run(sawyer_base)


def test_property_and_pricing(base: CopilotOutputs) -> None:
    s = base.summary
    assert s.units == 288 and s.occupied == 271
    assert s.avg_in_place_rent == 1_440 and s.avg_market_rent == 1_511
    assert s.loss_to_lease_pct == pytest.approx(0.047, abs=0.001)
    assert s.price_per_unit == 159_722
    assert 0.055 < s.going_in_cap < 0.060
    assert s.cap_at_ask is not None and s.cap_at_ask < s.going_in_cap
    assert s.discount_to_ask == pytest.approx(1 - 46.0 / 50.5)


def test_loan_is_ltv_bound(base: CopilotOutputs) -> None:
    loan = base.loan
    assert loan.binding == "LTV"
    assert loan.amount == 29_900_000
    assert loan.by_dscr > loan.amount and loan.by_debt_yield > loan.amount
    assert loan.annual_debt_service_io == pytest.approx(29_900_000 * 0.0575, abs=1)
    assert loan.annual_debt_service_amortizing > loan.annual_debt_service_io
    assert str(loan.maturity) == "2033-03-01" and str(loan.io_expiry) == "2029-03-01"
    assert loan.payoff < loan.amount


def test_sources_and_uses_foot(base: CopilotOutputs) -> None:
    su = base.sources_uses
    assert (
        su.total_uses
        == su.purchase_price
        + su.closing_costs
        + su.loan_fees
        + su.acquisition_fee
        + su.capital_budget
    )
    assert (
        su.capital_budget
        == su.renovation_budget + su.other_capital + su.construction_management_fee + su.contingency
    )
    assert su.equity == su.total_uses - su.loan
    assert su.lp_equity + su.gp_equity == pytest.approx(su.equity, abs=1)
    assert su.renovation_budget == 120 * 9_500


def test_annual_table_foots(base: CopilotOutputs) -> None:
    for row in base.annual:
        assert row.net_rental == pytest.approx(
            row.gpr
            + row.loss_to_lease
            + row.renovation_premium
            + row.renovation_vacancy
            + row.vacancy
            + row.concessions
            + row.non_revenue
            + row.bad_debt,
            abs=len(base.annual) + 8,
        )
        assert row.egi == pytest.approx(row.net_rental + row.other_income, abs=2)
        assert row.opex == pytest.approx(row.controllable + row.taxes + row.management_fee, abs=3)
        assert row.noi == pytest.approx(row.egi - row.opex, abs=2)
        assert row.cash_flow == pytest.approx(
            row.noi - row.reserves - row.debt_service - row.asset_management_fee, abs=4
        )
    assert len(base.annual) == 5
    assert [round(r.dscr or 0, 2) for r in base.annual][0] == pytest.approx(
        base.summary.dscr_year1 or 0, abs=0.01
    )
    # Interest only for three years, then amortizing debt service steps up.
    assert base.annual[2].debt_service < base.annual[3].debt_service
    assert base.annual[0].debt_service == pytest.approx(base.loan.annual_debt_service_io, abs=2)


def test_returns_clear_threshold_and_ask_does_not(
    base: CopilotOutputs, sawyer_base: CopilotInputs
) -> None:
    r = base.returns
    assert r.levered_irr is not None and 0.13 < r.levered_irr < 0.16
    assert r.unlevered_irr is not None and r.unlevered_irr < r.levered_irr
    assert 1.8 < r.equity_multiple < 2.0
    assert r.lp_irr is not None and r.lp_irr < r.levered_irr
    assert r.gp_irr is not None and r.gp_irr > r.levered_irr  # promote
    assert r.promote > 0
    at_ask = run(at_price(sawyer_base, 50_500_000)).summary
    assert at_ask.levered_irr is not None and at_ask.levered_irr < 0.10
    assert at_ask.going_in_cap == pytest.approx(base.summary.cap_at_ask or 0, abs=1e-6)


def test_waterfall_tiers_foot_to_distributions(base: CopilotOutputs) -> None:
    total_to_equity = base.sources_uses.equity * base.returns.equity_multiple
    assert sum(base.waterfall.tier_totals) == pytest.approx(total_to_equity, rel=1e-6)
    assert len(base.waterfall.tier_labels) == 3
    assert sum(base.waterfall.lp_by_year) + sum(base.waterfall.gp_by_year) == pytest.approx(
        total_to_equity, abs=10
    )


def test_exit_on_forward_noi(base: CopilotOutputs) -> None:
    e = base.exit
    assert e.gross_value == pytest.approx(e.forward_noi / 0.055, rel=1e-6)
    assert e.net_proceeds == e.gross_value - e.sale_costs
    assert e.net_after_debt == e.net_proceeds - e.loan_payoff
    assert e.forward_noi > base.annual[-1].noi


def test_cases_move_in_the_expected_direction(base: CopilotOutputs) -> None:
    downside = run(load_case("Downside"))
    lender = run(load_case("Lender"))
    assert downside.loan.binding == "Held" and downside.loan.amount == base.loan.amount
    assert downside.summary.levered_irr is not None and downside.summary.levered_irr < 0.05
    assert downside.summary.min_dscr is not None and downside.summary.min_dscr >= 1.25
    assert lender.loan.amount < base.loan.amount and lender.summary.ltv == pytest.approx(0.55)
    assert lender.summary.dscr_year1 is not None and lender.summary.dscr_year1 > (
        base.summary.dscr_year1 or 0
    )


def test_stress_table_holds_loan_and_covenant(
    sawyer_base: CopilotInputs, base: CopilotOutputs
) -> None:
    rows = stress_table(sawyer_base)
    assert [r.label for r in rows][:2] == ["Exit cap 6.25%", "Occupancy 90%"]
    assert "Renovation premium $75" in [r.label for r in rows]
    for row in rows:
        assert row.levered_irr is not None and row.levered_irr < (base.returns.levered_irr or 0)
        assert row.min_dscr is not None and row.holds
    exit_cap = rows[0]
    assert exit_cap.dscr_year1 == pytest.approx(
        base.summary.dscr_year1 or 0, abs=1e-9
    )  # exit stress leaves debt service alone


def assert_matches(actual: Any, expected: Any, path: str) -> None:
    if isinstance(expected, dict):
        assert isinstance(actual, dict), path
        assert actual.keys() == expected.keys(), path
        for key in expected:
            assert_matches(actual[key], expected[key], f"{path}.{key}")
    elif isinstance(expected, list):
        assert isinstance(actual, list) and len(actual) == len(expected), path
        for i, (a, e) in enumerate(zip(actual, expected, strict=True)):
            assert_matches(a, e, f"{path}[{i}]")
    elif isinstance(expected, float):
        assert actual == pytest.approx(expected, abs=1e-6), path
    else:
        assert actual == expected, path


def test_base_matches_fixture(base: CopilotOutputs) -> None:
    expected = json.loads(FIXTURE.read_text(encoding="utf-8"))
    assert_matches(base.model_dump(mode="json"), expected, "outputs")


def test_export_workbook_has_live_formulas(
    sawyer_base: CopilotInputs, base: CopilotOutputs
) -> None:
    wb = load_workbook(io.BytesIO(export_workbook(sawyer_base, base, "Base")))
    assert wb.sheetnames == ["Inputs", "Summary", "Pro forma", "Annual", "Waterfall"]
    summary = wb["Summary"]
    assert str(summary["B6"].value).startswith("=Annual!")  # year 1 NOI
    assert str(summary["B7"].value).startswith("=")  # going-in cap
    annual = wb["Annual"]
    assert any(
        "SUMIF(" in str(c.value) for row in annual.iter_rows(min_row=5, max_row=8) for c in row
    )
    pro_forma = wb["Pro forma"]
    header = [c.value for c in pro_forma[4]]
    noi_col = header.index("NOI") + 1
    assert str(pro_forma.cell(row=5, column=noi_col).value).startswith("=")
