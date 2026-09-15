"""Deal documents as text: the offering memorandum, the rent roll and the T-12.

PDFs are read page by page; spreadsheets keep their cell grid for the structured parsers and are
also rendered as numbered lines so a citation can point at a row. Files are limited by size and
type before anything is parsed.
"""

from __future__ import annotations

import csv
import datetime
import io
from pathlib import PurePosixPath
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader

Kind = Literal["om", "rent_roll", "t12"]
KIND_LABELS: dict[str, str] = {"om": "OM", "rent_roll": "Rent roll", "t12": "T-12"}
MAX_BYTES = 8 * 1024 * 1024
ALLOWED_SUFFIXES = {
    "om": {".pdf"},
    "rent_roll": {".xlsx", ".csv", ".pdf"},
    "t12": {".xlsx", ".csv", ".pdf"},
}
MAX_PAGES = 120
MAX_ROWS = 5_000

Cell = str | float | int | datetime.datetime | datetime.date | None


class DocumentError(ValueError):
    """A plain-language reason the file could not be used."""


class Page(BaseModel):
    number: int
    text: str


class Document(BaseModel):
    kind: Kind
    filename: str
    pages: list[Page] = Field(default_factory=list)
    # Spreadsheet cells by row, for the structured parsers. Empty for PDFs.
    rows: list[list[Cell]] = Field(default_factory=list)

    @property
    def label(self) -> str:
        return KIND_LABELS[self.kind]

    def text(self) -> str:
        return "\n".join(p.text for p in self.pages)

    def page(self, number: int) -> Page | None:
        for p in self.pages:
            if p.number == number:
                return p
        return None


def read_document(kind: Kind, filename: str, data: bytes) -> Document:
    suffix = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES[kind]:
        allowed = ", ".join(sorted(ALLOWED_SUFFIXES[kind]))
        raise DocumentError(f"{KIND_LABELS[kind]}: {filename} is not accepted; use {allowed}.")
    if len(data) > MAX_BYTES:
        raise DocumentError(
            f"{KIND_LABELS[kind]}: {filename} is over {MAX_BYTES // (1024 * 1024)} MB."
        )
    if len(data) == 0:
        raise DocumentError(f"{KIND_LABELS[kind]}: {filename} is empty.")
    if suffix == ".pdf":
        return _read_pdf(kind, filename, data)
    if suffix == ".xlsx":
        return _read_xlsx(kind, filename, data)
    return _read_csv(kind, filename, data)


def _read_pdf(kind: Kind, filename: str, data: bytes) -> Document:
    if not data.startswith(b"%PDF"):
        raise DocumentError(f"{KIND_LABELS[kind]}: {filename} is not a PDF.")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError(f"{KIND_LABELS[kind]}: {filename} is password protected.")
        pages = [
            Page(number=i, text=(page.extract_text() or "").strip())
            for i, page in enumerate(reader.pages[:MAX_PAGES], start=1)
        ]
    except DocumentError:
        raise
    except Exception as exc:  # pypdf raises a family of parse errors
        raise DocumentError(f"{KIND_LABELS[kind]}: {filename} could not be read ({exc}).") from exc
    if not any(p.text for p in pages):
        raise DocumentError(
            f"{KIND_LABELS[kind]}: {filename} has no text layer; it is likely a scan. "
            "Export the document as text or run OCR before uploading."
        )
    return Document(kind=kind, filename=filename, pages=pages)


def _cell(value: object) -> Cell:
    if value is None or isinstance(value, (str, int, float, datetime.datetime, datetime.date)):
        return value
    return str(value)


def _read_xlsx(kind: Kind, filename: str, data: bytes) -> Document:
    if not data.startswith(b"PK"):
        raise DocumentError(f"{KIND_LABELS[kind]}: {filename} is not an Excel workbook.")
    try:
        # data_only returns cached formula results; a formula never calculated comes back None,
        # which is why the parsers sum the month columns themselves.
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentError(
            f"{KIND_LABELS[kind]}: {filename} could not be opened ({exc})."
        ) from exc
    # Row numbers match the spreadsheet (blank rows are kept as empty lists), so a citation such
    # as "row 19" can be checked against the file. Only the first sheet with data is parsed.
    rows: list[list[Cell]] = []
    pages: list[Page] = []
    for sheet_no, ws in enumerate(wb.worksheets, start=1):
        sheet_rows: list[list[Cell]] = []
        for row in ws.iter_rows(values_only=True):
            if len(sheet_rows) >= MAX_ROWS:
                break
            cells = [_cell(v) for v in row]
            sheet_rows.append(cells if any(c not in (None, "") for c in cells) else [])
        if not any(sheet_rows):
            continue
        if not rows:
            rows = sheet_rows
        pages.append(Page(number=sheet_no, text=_rows_as_text(sheet_rows, 1, ws.title)))
    if not rows:
        raise DocumentError(f"{KIND_LABELS[kind]}: {filename} has no data.")
    return Document(kind=kind, filename=filename, pages=pages, rows=rows)


def _read_csv(kind: Kind, filename: str, data: bytes) -> Document:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    rows: list[list[Cell]] = []
    for raw in csv.reader(io.StringIO(text)):
        if len(rows) >= MAX_ROWS:
            break
        cells: list[Cell] = [_number_or_text(c) for c in raw]
        if any(c not in (None, "") for c in cells):
            rows.append(cells)
    if not rows:
        raise DocumentError(f"{KIND_LABELS[kind]}: {filename} has no data.")
    return Document(
        kind=kind,
        filename=filename,
        pages=[Page(number=1, text=_rows_as_text(rows, 1, "csv"))],
        rows=rows,
    )


def _number_or_text(raw: str) -> Cell:
    text = raw.strip()
    if text == "":
        return None
    candidate = text.replace(",", "").replace("$", "")
    negative = candidate.startswith("(") and candidate.endswith(")")
    candidate = candidate.strip("()")
    try:
        number = float(candidate)
    except ValueError:
        return text
    return -number if negative else number


def _fmt(cell: Cell) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float):
        return f"{cell:,.0f}" if cell.is_integer() else f"{cell:,.2f}"
    if isinstance(cell, (datetime.datetime, datetime.date)):
        return cell.strftime("%Y-%m-%d")
    return str(cell)


def _rows_as_text(rows: list[list[Cell]], start: int, title: str) -> str:
    lines = [f"[{title}]"]
    for n, cells in enumerate(rows, start=start):
        if cells:
            lines.append(f"row {n}: " + " | ".join(_fmt(c) for c in cells).rstrip(" |"))
    return "\n".join(lines)
